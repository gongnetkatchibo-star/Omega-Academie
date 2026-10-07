"""Import des élèves depuis un fichier Excel.

Tout ou rien : le fichier est contrôlé ligne par ligne, et si une seule
ligne est en erreur, rien n'est importé — la liste des erreurs indique
quoi corriger. Les élèves déjà présents (même nom, même classe) sont
laissés de côté et signalés, jamais doublés."""

import datetime as dt
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.extensions import db

COLONNES = ["Nom complet", "Classe", "Date de naissance", "Genre", "Téléphone du parent"]
LIGNES_MAX = 2000
GENRES = {"m": "M", "masculin": "M", "garçon": "M", "garcon": "M", "h": "M",
          "f": "F", "féminin": "F", "feminin": "F", "fille": "F"}


def modele_excel(classes):
    """Fichier modèle : les colonnes attendues et, sur une 2e feuille,
    les noms de classes acceptés."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Élèves"
    feuille.append(COLONNES)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="0F5FA6")
    for lettre, largeur in zip("ABCDE", (34, 14, 20, 10, 24)):
        feuille.column_dimensions[lettre].width = largeur
    feuille.freeze_panes = "A2"

    aide = classeur.create_sheet("Classes acceptées")
    aide.append(["Classe"])
    aide["A1"].font = Font(bold=True)
    for classe in classes:
        aide.append([classe.nom])
    aide.column_dimensions["A"].width = 20

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def _texte(valeur):
    return "" if valeur is None else str(valeur).strip()


def _date(valeur):
    """Date Excel ou texte JJ/MM/AAAA. Retourne (date|None, erreur|None)."""
    if valeur is None or _texte(valeur) == "":
        return None, None
    if isinstance(valeur, dt.datetime):
        return valeur.date(), None
    if isinstance(valeur, dt.date):
        return valeur, None
    texte = _texte(valeur)
    for forme in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(texte, forme).date(), None
        except ValueError:
            continue
    return None, f"date de naissance « {texte} » illisible (attendu : JJ/MM/AAAA)"


def lire_fichier(fichier, classes):
    """Contrôle le fichier. Retourne (lignes valides, erreurs)."""
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception:
        return [], ["Ce fichier n'est pas un classeur Excel (.xlsx) lisible."]
    feuille = classeur.worksheets[0]
    rangees = feuille.iter_rows(values_only=True)

    entetes = [_texte(c).lower() for c in (next(rangees, None) or [])]
    attendues = [c.lower() for c in COLONNES]
    if entetes[:len(attendues)] != attendues:
        return [], ["La première ligne doit contenir exactement les colonnes du modèle : " + ", ".join(COLONNES) + "."]

    par_nom = {c.nom.strip().lower(): c for c in classes}
    lignes, erreurs = [], []
    aujourd_hui = dt.date.today()

    for numero, rangee in enumerate(rangees, start=2):
        cellules = list(rangee[:5]) + [None] * (5 - len(rangee[:5]))
        if all(_texte(c) == "" for c in cellules):
            continue
        if len(lignes) + len(erreurs) >= LIGNES_MAX:
            erreurs.append(f"Le fichier dépasse {LIGNES_MAX} élèves : découpe-le en plusieurs fichiers.")
            break
        nom, classe_nom, naissance, genre, telephone = cellules
        nom, classe_nom, genre, telephone = _texte(nom), _texte(classe_nom), _texte(genre), _texte(telephone)
        problemes = []
        if not nom:
            problemes.append("nom manquant")
        elif len(nom) > 120:
            problemes.append("nom trop long")
        classe = par_nom.get(classe_nom.lower())
        if not classe_nom:
            problemes.append("classe manquante")
        elif classe is None:
            problemes.append(f"classe « {classe_nom} » inconnue pour cette année")
        date_naissance, erreur_date = _date(naissance)
        if erreur_date:
            problemes.append(erreur_date)
        elif date_naissance and not dt.date(1990, 1, 1) <= date_naissance <= aujourd_hui:
            problemes.append("date de naissance invraisemblable")
        sexe = GENRES.get(genre.lower()) if genre else None
        if genre and sexe is None:
            problemes.append(f"genre « {genre} » inconnu (M ou F)")
        if len(telephone) > 30:
            problemes.append("téléphone trop long")

        if problemes:
            erreurs.append(f"Ligne {numero} : " + " ; ".join(problemes) + ".")
        else:
            lignes.append({"nom": nom, "classe": classe, "date_naissance": date_naissance,
                           "sexe": sexe, "telephone": telephone or None})
    if not lignes and not erreurs:
        erreurs.append("Le fichier ne contient aucun élève.")
    return lignes, erreurs


def importer(lignes):
    """Crée les élèves. Retourne (nombre créés, nombre déjà présents)."""
    from app.models.eleve import Eleve
    from app.models.historique import HistoriqueScolaire

    annee = Eleve.annee_scolaire_courante()
    existants = {
        (e.nom_complet.strip().lower(), e.classe_id)
        for e in Eleve.query.filter(Eleve.classe_id.in_({l["classe"].id for l in lignes})).all()
    }
    crees = deja = 0
    for ligne in lignes:
        classe = ligne["classe"]
        cle = (ligne["nom"].lower(), classe.id)
        if cle in existants:
            deja += 1
            continue
        existants.add(cle)
        eleve = Eleve(
            matricule=Eleve.generer_matricule(classe), nom_complet=ligne["nom"], classe_id=classe.id,
            date_naissance=ligne["date_naissance"], sexe=ligne["sexe"], telephone_parent=ligne["telephone"],
        )
        db.session.add(eleve)
        db.session.flush()
        eleve.actualiser_statut_dossier()
        db.session.add(HistoriqueScolaire(eleve_id=eleve.id, classe_id=classe.id, annee_scolaire=annee, resultat="en_cours"))
        crees += 1
    return crees, deja
