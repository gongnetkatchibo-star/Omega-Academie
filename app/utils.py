from functools import wraps

from flask import abort
from flask_login import current_user


def roles_required(*roles, module=None):
    """Restreint une vue aux utilisateurs dont le rôle figure dans
    `roles` — sauf si `module` est précisé et qu'une permission
    explicite existe pour ce couple (rôle, module) depuis la matrice de
    permissions de l'espace développeur ; dans ce cas, elle prend le
    dessus sur `roles` (sept. 2026).

    Le rôle « developpeur » a toujours accès complet à toute vue protégée
    par ce décorateur, quels que soient les rôles listés — c'est la seule
    exception volontaire (décision du fondateur)."""

    def decorateur(f):
        @wraps(f)
        def enveloppe(*args, **kwargs):
            if not current_user.is_authenticated:
                abort(403)
            if module:
                from app.services.permissions import role_a_acces
                autorise = role_a_acces(current_user.role, module, roles)
            else:
                autorise = current_user.role == "developpeur" or current_user.role in roles
            if not autorise:
                abort(403)
            return f(*args, **kwargs)

        return enveloppe

    return decorateur


def logo_officiel_data_uri():
    """Encode le logo officiel en data-URI — xhtml2pdf gère mal les
    chemins de fichiers relatifs, l'encodage direct est fiable partout
    (sept. 2026)."""
    import base64
    import os

    chemin = os.path.join(os.path.dirname(__file__), "static", "images", "logo-csoa-officiel.png")
    with open(chemin, "rb") as f:
        contenu = base64.b64encode(f.read()).decode("ascii")
    return f"data:image/png;base64,{contenu}"


def filigrane_data_uri():
    """Filigrane officiel (Ω + livre, sans texte), extrait du document de
    référence de l'école — déjà pâli, utilisé tel quel (sept. 2026)."""
    import base64
    import os

    chemin = os.path.join(os.path.dirname(__file__), "static", "images", "filigrane-csoa.png")
    with open(chemin, "rb") as f:
        contenu = base64.b64encode(f.read()).decode("ascii")
    return f"data:image/png;base64,{contenu}"


def html_vers_pdf(html):
    """Convertit un fragment HTML (déjà rendu par un template) en PDF,
    avec le filigrane officiel placé exactement comme sur le document de
    référence de l'école : sur une page A4 (595 x 842 pt), l'image occupe
    x 71 → 524 pt et y 260 → 631 pt depuis le haut (453 x 371 pt).
    Pour un autre format de page, la même proportion est conservée."""
    import re
    from io import BytesIO
    from xhtml2pdf import pisa

    largeur_page, hauteur_page = 595.0, 842.0
    taille = re.search(r"@page\s*\{[^}]*?size:\s*([\d.]+)pt\s+([\d.]+)pt", html)
    if taille:
        largeur_page, hauteur_page = float(taille.group(1)), float(taille.group(2))

    echelle = min(largeur_page / 595.0, hauteur_page / 842.0)
    largeur_img, hauteur_img = 453.0 * echelle, 371.0 * echelle
    x = (largeur_page - largeur_img) / 2
    haut_depuis_le_haut = 260.0 / 842.0 * hauteur_page
    y_depuis_le_bas = hauteur_page - haut_depuis_le_haut - hauteur_img

    proprietes = (
        f'background-image: url("{filigrane_data_uri()}"); '
        f"background-object-position: {x:.0f}pt {y_depuis_le_bas:.0f}pt; "
        f"background-width: {largeur_img:.0f}pt; background-height: {hauteur_img:.0f}pt;"
    )

    # xhtml2pdf ne fusionne pas deux règles @page : on complète celle du
    # modèle s'il en a une, sinon on en crée une.
    page_existante = re.search(r"@page\s*\{", html)
    if page_existante:
        html = html[:page_existante.end()] + proprietes + html[page_existante.end():]
    else:
        html = re.sub(r"(<head[^>]*>)", r"\1<style>@page {" + proprietes + "}</style>", html, count=1)

    tampon = BytesIO()
    pisa.CreatePDF(src=html, dest=tampon, encoding="utf-8")
    tampon.seek(0)
    return tampon.read()


def export_csv(entetes, lignes, nom_fichier):
    """Génère une réponse Flask téléchargeable au format CSV.
    `lignes` est une liste de tuples/listes, dans le même ordre que
    `entetes`."""
    import csv
    from io import StringIO
    from flask import Response

    tampon = StringIO()
    writer = csv.writer(tampon, delimiter=";")
    writer.writerow(entetes)
    writer.writerows(lignes)
    return Response(
        tampon.getvalue().encode("utf-8-sig"),  # BOM pour un bon affichage des accents dans Excel
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={nom_fichier}.csv"},
    )


def export_xlsx(entetes, lignes, nom_fichier, titre_feuille="Export"):
    """Génère une réponse Flask téléchargeable au format Excel (.xlsx)."""
    from io import BytesIO
    from flask import Response
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titre_feuille[:31]  # limite Excel pour le nom d'un onglet

    ws.append(list(entetes))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="00387B", end_color="00387B", fill_type="solid")
    for ligne in lignes:
        ws.append(list(ligne))
    for col in ws.columns:
        longueur = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(longueur + 2, 45)

    tampon = BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return Response(
        tampon.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nom_fichier}.xlsx"},
    )


def export_pdf_liste(titre, sous_titre, entetes, lignes, nom_fichier):
    """Génère un PDF imprimable listant des lignes sous forme de tableau
    simple — utilisé pour les exports de listes (enseignants, élèves…).
    Porte désormais l'en-tête officiel CSOA, comme tout document imprimé
    de l'établissement (sept. 2026)."""
    from flask import make_response, render_template
    from app.services.documents_officiels import contexte_entete_officiel

    html = render_template(
        "exports/liste_pdf.html", titre=titre, sous_titre=sous_titre,
        entetes=entetes, lignes=lignes, etablissement="Omega Académie",
        **contexte_entete_officiel(),
    )
    reponse = make_response(html_vers_pdf(html))
    reponse.headers["Content-Type"] = "application/pdf"
    reponse.headers["Content-Disposition"] = f"attachment; filename={nom_fichier}.pdf"
    return reponse


# Extensions autorisées pour les fichiers envoyés par les utilisateurs —
# jamais d'exécutable ni de script, quel que soit le module (sept. 2026).
EXTENSIONS_DOCUMENTS = {"pdf", "doc", "docx", "ppt", "pptx", "xls", "xlsx", "txt", "csv"}
EXTENSIONS_IMAGES = {"jpg", "jpeg", "png", "gif", "webp"}
EXTENSIONS_VIDEOS = {"mp4", "webm", "mov"}
EXTENSIONS_BIBLIOTHEQUE = EXTENSIONS_DOCUMENTS | EXTENSIONS_IMAGES | EXTENSIONS_VIDEOS
EXTENSIONS_PIECE_JOINTE = EXTENSIONS_DOCUMENTS | EXTENSIONS_IMAGES


def extension_autorisee(nom_fichier, extensions_autorisees):
    if "." not in nom_fichier:
        return False
    extension = nom_fichier.rsplit(".", 1)[1].lower()
    return extension in extensions_autorisees


def normaliser_numero_tchad(saisie):
    """Valide et normalise un numéro mobile tchadien.

    Accepte avec ou sans indicatif (+235, 00235, ou rien), espaces
    ignorés. Un mobile tchadien a 8 chiffres et commence par 6 ou 9
    (les numéros en 22 sont des lignes fixes, exclues ici puisqu'on veut
    pouvoir joindre la personne sur mobile). Retourne le numéro au
    format +235XXXXXXXX, ou None si invalide.

    Remarque : depuis la portabilité des numéros (2020), le premier
    chiffre ne garantit plus l'opérateur d'origine à 100% — c'est pour
    ça que l'opérateur est choisi par la personne elle-même dans le
    formulaire, plutôt que deviné automatiquement à partir du numéro."""
    import re

    chiffres = re.sub(r"[^0-9]", "", saisie)

    if chiffres.startswith("00235"):
        chiffres = chiffres[5:]
    elif chiffres.startswith("235"):
        chiffres = chiffres[3:]

    if len(chiffres) != 8 or chiffres[0] not in "69":
        return None

    return f"+235{chiffres}"
