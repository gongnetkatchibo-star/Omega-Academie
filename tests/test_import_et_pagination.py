"""Import Excel des élèves, pagination et recherche."""

import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import connecter


def _classeur(lignes, entetes=("Nom complet", "Classe", "Date de naissance", "Genre", "Téléphone du parent")):
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(list(entetes))
    for ligne in lignes:
        feuille.append(list(ligne))
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def _importer(client, tampon, nom="eleves.xlsx"):
    return client.post("/eleves/importer", data={"fichier": (tampon, nom)},
                       content_type="multipart/form-data", follow_redirects=True).get_data(as_text=True)


@pytest.fixture
def secretaire(client, creer_utilisateur, creer_classe):
    creer_utilisateur("Sec", "s@t.com", "secretaire")
    creer_classe(nom="CP1", niveau=1)
    creer_classe(nom="6e", niveau=7)
    connecter(client, "s@t.com")


def test_import_cree_les_eleves_avec_matricule_et_historique(client, db, secretaire):
    from app.models.eleve import Eleve
    from app.models.historique import HistoriqueScolaire

    page = _importer(client, _classeur([
        ("Awa Mahamat", "CP1", datetime(2019, 4, 2), "F", "+23566000001"),
        ("Brahim Ali", "cp1", "15/09/2018", "garçon", None),
        ("Clé Deby", "6e", None, None, None),
        (None, None, None, None, None),
    ]))
    assert "3 élève(s) importé(s)." in page
    db.session.expire_all()
    eleves = {e.nom_complet: e for e in Eleve.query.all()}
    assert eleves["Awa Mahamat"].matricule == "ET26-CP1-001" and eleves["Brahim Ali"].matricule == "ET26-CP1-002"
    assert eleves["Clé Deby"].matricule == "ET26-6E-001"
    assert eleves["Awa Mahamat"].date_naissance == date(2019, 4, 2) and eleves["Awa Mahamat"].sexe == "F"
    assert eleves["Brahim Ali"].date_naissance == date(2018, 9, 15) and eleves["Brahim Ali"].sexe == "M"
    assert eleves["Awa Mahamat"].telephone_parent == "+23566000001"
    assert HistoriqueScolaire.query.count() == 3


def test_une_erreur_et_rien_n_est_importe(client, db, secretaire):
    from app.models.eleve import Eleve

    page = _importer(client, _classeur([
        ("Bonne Ligne", "CP1", None, "F", None),
        ("", "CP1", None, None, None),
        ("Mauvaise Classe", "Terminale", "32/13/2020", "X", None),
    ]))
    assert "Rien n&#39;a été importé" in page or "Rien n'a été importé" in page
    assert "Ligne 3 : nom manquant." in page
    assert "Ligne 4 :" in page and "Terminale" in page and "illisible" in page and "genre" in page
    db.session.expire_all()
    assert Eleve.query.count() == 0


def test_reimport_ne_double_pas(client, db, secretaire):
    from app.models.eleve import Eleve

    lignes = [("Awa Mahamat", "CP1", None, "F", None)]
    _importer(client, _classeur(lignes))
    page = _importer(client, _classeur(lignes + [("Nouveau Venu", "CP1", None, "M", None)]))
    assert "1 élève(s) importé(s)." in page and "1 déjà présent(s)" in page
    db.session.expire_all()
    assert Eleve.query.count() == 2
    assert sorted(e.matricule for e in Eleve.query.all()) == ["ET26-CP1-001", "ET26-CP1-002"]


def test_fichiers_refuses(client, db, secretaire):
    assert "format Excel" in _importer(client, io.BytesIO(b"a,b"), nom="eleves.csv")
    assert "lisible" in _importer(client, io.BytesIO(b"pas un excel"))
    assert "colonnes du modèle" in _importer(client, _classeur([("x", "CP1")], entetes=("Nom", "Classe")))
    assert "aucun élève" in _importer(client, _classeur([]))


def test_modele_telechargeable(client, secretaire):
    r = client.get("/eleves/importer/modele")
    assert r.status_code == 200
    classeur = load_workbook(io.BytesIO(r.get_data()))
    assert [c.value for c in classeur.worksheets[0][1]] == ["Nom complet", "Classe", "Date de naissance", "Genre", "Téléphone du parent"]
    assert [r[0].value for r in classeur.worksheets[1].iter_rows(min_row=2)] == ["CP1", "6e"]


def test_import_reserve_a_la_gestion(client, creer_utilisateur):
    creer_utilisateur("Prof", "e@t.com", "enseignant")
    connecter(client, "e@t.com")
    assert client.get("/eleves/importer").status_code == 403


def test_liste_des_eleves_paginee_et_recherche(client, db, secretaire, creer_eleve):
    from app.models.classe import Classe

    classe = Classe.query.filter_by(nom="CP1").one()
    for i in range(120):
        creer_eleve(f"Eleve {i:03d}", classe, matricule=f"ET26-CP1-{i:03d}")
    creer_eleve("Zara Unique", classe, matricule="ET26-CP1-999")

    page1 = client.get("/eleves/").get_data(as_text=True)
    assert "Eleve 000" in page1 and "Eleve 049" in page1 and "Eleve 050" not in page1
    assert "1–50 sur 121" in page1 and "page 1/3" in page1
    page3 = client.get("/eleves/?page=3").get_data(as_text=True)
    assert "Zara Unique" in page3 and "101–121 sur 121" in page3
    assert "Eleve 000" in client.get("/eleves/?page=0").get_data(as_text=True)
    assert "Zara Unique" in client.get("/eleves/?page=99").get_data(as_text=True)

    trouve = client.get("/eleves/?q=zara").get_data(as_text=True)
    assert "Zara Unique" in trouve and "Eleve 000" not in trouve
    assert "Zara Unique" in client.get("/eleves/?q=CP1-999").get_data(as_text=True)
    assert "Aucun élève ne correspond" in client.get("/eleves/?q=introuvable").get_data(as_text=True)
    # Le filtre est conservé d'une page à l'autre.
    assert "q=Eleve" in client.get("/eleves/?q=Eleve").get_data(as_text=True)


def test_recherche_globale(client, db, secretaire, creer_eleve, creer_utilisateur):
    from app.models.classe import Classe

    creer_eleve("Zara Unique", Classe.query.filter_by(nom="CP1").one(), matricule="ET26-CP1-777")
    creer_utilisateur("Zara Comptable", "zc@t.com", "comptable")
    page = client.get("/recherche?q=zara").get_data(as_text=True)
    assert "Zara Unique" in page and "ET26-CP1-777" in page
    assert "zc@t.com" not in page  # le secrétariat ne gère pas les comptes
    assert "Aucun résultat" in client.get("/recherche?q=xyzxyz").get_data(as_text=True)
    assert 'action="/recherche"' in client.get("/").get_data(as_text=True)


def test_recherche_refusee_aux_parents(client, creer_utilisateur):
    creer_utilisateur("Parent", "p@t.com", "parent")
    connecter(client, "p@t.com")
    assert client.get("/recherche?q=awa").status_code == 403
    assert 'action="/recherche"' not in client.get("/").get_data(as_text=True)


def test_menu_regroupe(client, secretaire):
    page = client.get("/").get_data(as_text=True)
    for libelle in ("Scolarité", "Pédagogie", "Administration", "Documents officiels"):
        assert libelle in page
    assert page.index("Tableau de bord") < page.index("Scolarité")
    assert client.get("/documents/").status_code == 200


def test_index_crees_au_demarrage(app, db):
    from sqlalchemy import inspect

    noms = {i["name"] for i in inspect(db.engine).get_indexes("eleves")}
    assert {"ix_eleves_classe", "ix_eleves_nom"} <= noms
